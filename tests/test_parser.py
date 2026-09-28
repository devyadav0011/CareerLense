import os
import tempfile
import pytest
from app.services.parser import ResumeParser, extract_text_from_pdf, extract_text_from_docx

def test_parse_text(sample_raw_text):
    result = ResumeParser.parse_text(sample_raw_text)
    assert 'personal_info' in result
    assert result['personal_info']['full_name'] == 'Jane Doe'
    assert result['personal_info']['email'] == 'jane.doe@example.com'
    assert '+1 (555) 234-5678' in result['personal_info']['phone']
    assert 'linkedin.com/in/janedoe' in result['personal_info']['linkedin']
    assert 'github.com/janedoe' in result['personal_info']['github']

    # Skills taxonomy
    skills = result['skills']
    assert 'Python' in skills['programming_languages']
    assert 'Flask' in skills['frameworks']
    assert 'PostgreSQL' in skills['databases']
    assert 'Docker' in skills['cloud']
    assert 'Git' in skills['tools']

    # Sections
    assert 'education' in result['detected_sections']
    assert 'experience' in result['detected_sections']
    assert 'skills' in result['detected_sections']

def test_pdf_parsing(sample_raw_text):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tf:
        pdf_path = tf.name

    try:
        doc = SimpleDocTemplate(pdf_path, pagesize=A4)
        styles = getSampleStyleSheet()
        story = [Paragraph("Jane Doe Software Engineer", styles['Normal']), Paragraph("Email: jane@example.com", styles['Normal'])]
        doc.build(story)

        extracted = extract_text_from_pdf(pdf_path)
        assert "Jane Doe" in extracted
        assert "jane@example.com" in extracted
    finally:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)

def test_docx_parsing():
    import docx

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tf:
        docx_path = tf.name

    try:
        doc = docx.Document()
        doc.add_paragraph("Devan Yadav")
        doc.add_paragraph("Full-Stack Developer with Python experience")
        doc.save(docx_path)

        extracted = extract_text_from_docx(docx_path)
        assert "Devan Yadav" in extracted
        assert "Full-Stack Developer" in extracted
    finally:
        if os.path.exists(docx_path):
            os.remove(docx_path)

def test_empty_pdf_error():
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tf:
        empty_pdf = tf.name

    try:
        with open(empty_pdf, 'wb') as f:
            f.write(b"")
        with pytest.raises(Exception):
            extract_text_from_pdf(empty_pdf)
    finally:
        if os.path.exists(empty_pdf):
            os.remove(empty_pdf)
